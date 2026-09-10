#!/usr/bin/env python3
"""Create a bounded, local-only inventory of modeling project inputs."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path


EXCLUDED_DIRS = {
    ".git", ".idea", ".vscode", "__pycache__", "node_modules", ".venv", "venv",
    "code", "figures", "paper", "reports", "results", "_tmp",
}
INTERESTING_SUFFIXES = {
    ".csv", ".tsv", ".xlsx", ".xls", ".pdf", ".docx", ".txt", ".md",
    ".json", ".zip", ".rar", ".7z", ".png", ".jpg", ".jpeg", ".tif", ".tiff",
}


def compact(value: object, limit: int = 80) -> str:
    text = " ".join(str(value if value is not None else "").split())
    return text if len(text) <= limit else text[: limit - 1] + "…"


def human_size(size: int) -> str:
    value = float(size)
    for unit in ("B", "KB", "MB", "GB"):
        if value < 1024 or unit == "GB":
            return f"{value:.1f} {unit}"
        value /= 1024
    return f"{size} B"


def digest(path: Path, limit_bytes: int) -> str:
    if path.stat().st_size > limit_bytes:
        return "跳过（大文件）"
    sha = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            sha.update(block)
    return sha.hexdigest()[:16]


def open_text(path: Path):
    last_error: Exception | None = None
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            handle = path.open(encoding=encoding, newline="")
            handle.read(4096)
            handle.seek(0)
            return handle, encoding
        except UnicodeError as exc:
            last_error = exc
            try:
                handle.close()
            except UnboundLocalError:
                pass
    raise UnicodeError(f"无法识别文本编码: {last_error}")


def inspect_delimited(path: Path, sample_rows: int, large: bool) -> dict:
    handle, encoding = open_text(path)
    with handle:
        sample = handle.read(8192)
        handle.seek(0)
        fallback = "\t" if path.suffix.lower() == ".tsv" else ","
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
            delimiter = dialect.delimiter
        except csv.Error:
            delimiter = fallback
        reader = csv.reader(handle, delimiter=delimiter)
        preview = []
        for _, row in zip(range(sample_rows + 1), reader):
            preview.append([compact(cell) for cell in row])
        row_count = None
        if not large:
            row_count = len(preview)
            row_count += sum(1 for _ in reader)
    return {
        "encoding": encoding,
        "delimiter": repr(delimiter),
        "rows": row_count,
        "columns": len(preview[0]) if preview else 0,
        "header": preview[0] if preview else [],
        "preview": preview[1:],
    }


def inspect_xlsx(path: Path, sample_rows: int) -> dict:
    try:
        import openpyxl
    except ImportError:
        return {"warning": "缺少 openpyxl，仅记录文件信息"}
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    try:
        for worksheet in workbook.worksheets:
            preview = []
            for row in worksheet.iter_rows(min_row=1, max_row=sample_rows + 1, values_only=True):
                preview.append([compact(cell) for cell in row])
            sheets.append({
                "name": worksheet.title,
                "rows": worksheet.max_row,
                "columns": worksheet.max_column,
                "header": preview[0] if preview else [],
                "preview": preview[1:],
            })
    finally:
        workbook.close()
    return {"sheets": sheets}


def inspect_pdf(path: Path) -> dict:
    try:
        from pypdf import PdfReader
    except ImportError:
        return {"warning": "缺少 pypdf，仅记录文件信息"}
    try:
        return {"pages": len(PdfReader(str(path)).pages)}
    except Exception as exc:
        return {"warning": f"PDF 元数据读取失败：{compact(exc)}"}


def discover(root: Path) -> list[Path]:
    files = []
    for path in root.rglob("*"):
        relative = path.relative_to(root)
        if any(part in EXCLUDED_DIRS for part in relative.parts):
            continue
        if path.is_file() and path.suffix.lower() in INTERESTING_SUFFIXES:
            files.append(path)
    return sorted(files, key=lambda item: item.relative_to(root).as_posix().lower())


def render(root: Path, records: list[dict], threshold_mb: int, sample_rows: int) -> str:
    lines = [
        "# 输入附件清单",
        "",
        f"生成时间：{datetime.now(timezone.utc).isoformat(timespec='seconds')}",
        f"项目目录：`{root.name}`",
        f"大文件阈值：{threshold_mb} MB；表格预览行数：{sample_rows}",
        "",
        "> 本报告仅在本地读取附件；不执行宏、不上传文件。样例值会截断，正式计算仍须使用完整数据。",
        "",
        "## 文件概览",
        "",
        "| 相对路径 | 类型 | 大小 | 大文件 | SHA-256（前16位） |",
        "| --- | --- | ---: | --- | --- |",
    ]
    for record in records:
        lines.append(
            f"| `{record['path']}` | {record['suffix'] or '无扩展名'} | {record['size_human']} | "
            f"{'是' if record['large'] else '否'} | `{record['sha256']}` |"
        )
    if not records:
        lines.append("| _未发现支持的附件_ |  |  |  |  |")
    lines.extend(["", "## 结构预览", ""])
    for record in records:
        details = record.get("details", {})
        if not details:
            continue
        lines.append(f"### `{record['path']}`")
        if "warning" in details:
            lines.extend(["", f"- {details['warning']}", ""])
            continue
        if record["suffix"] in (".csv", ".tsv"):
            row_text = details["rows"] if details["rows"] is not None else "未全量计数（大文件）"
            lines.extend([
                "",
                f"- 编码：`{details['encoding']}`；分隔符：`{details['delimiter']}`",
                f"- 规模：{row_text} 行，{details['columns']} 列",
                f"- 字段：{', '.join(f'`{compact(x)}`' for x in details['header']) or '未识别'}",
            ])
            if details["preview"]:
                lines.append(f"- 前 {len(details['preview'])} 行截断预览：`{compact(details['preview'], 300)}`")
            lines.append("")
        elif record["suffix"] == ".xlsx":
            lines.append("")
            for sheet in details.get("sheets", []):
                lines.append(
                    f"- 工作表 `{compact(sheet['name'])}`：{sheet['rows']} 行 × {sheet['columns']} 列；"
                    f"字段：{', '.join(f'`{compact(x)}`' for x in sheet['header']) or '未识别'}"
                )
            lines.append("")
        elif record["suffix"] == ".pdf":
            lines.extend(["", f"- 页数：{details.get('pages', '未知')}", ""])
    lines.extend([
        "## 后续处理提示",
        "",
        "- 标记为大文件的 CSV/TSV 应使用分块读取、抽样探查和全量聚合，不要一次性载入内存。",
        "- XLSX 先按工作表和字段确认口径；公式以缓存值读取，脚本不会执行宏。",
        "- 任何抽样只用于理解结构，不得冒充全量统计或最终评价结果。",
    ])
    return "\n".join(lines) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description="生成数学建模项目的本地附件清单与有界预览")
    parser.add_argument("project_root", type=Path, help="项目根目录")
    parser.add_argument("--out", type=Path, help="输出 Markdown，默认 reports/INPUT_INVENTORY.md")
    parser.add_argument("--sample-rows", type=int, default=5)
    parser.add_argument("--large-file-mb", type=int, default=50)
    args = parser.parse_args()

    root = args.project_root.resolve()
    if not root.is_dir():
        parser.error(f"项目目录不存在：{root}")
    sample_rows = max(1, min(args.sample_rows, 20))
    threshold_mb = max(1, args.large_file_mb)
    threshold_bytes = threshold_mb * 1024 * 1024
    records = []
    for path in discover(root):
        size = path.stat().st_size
        suffix = path.suffix.lower()
        record = {
            "path": path.relative_to(root).as_posix(),
            "suffix": suffix,
            "size_human": human_size(size),
            "large": size >= threshold_bytes,
            "sha256": digest(path, threshold_bytes),
        }
        try:
            if suffix in (".csv", ".tsv"):
                record["details"] = inspect_delimited(path, sample_rows, record["large"])
            elif suffix == ".xlsx":
                record["details"] = inspect_xlsx(path, sample_rows)
            elif suffix == ".pdf":
                record["details"] = inspect_pdf(path)
        except Exception as exc:
            record["details"] = {"warning": f"结构预检失败：{compact(exc)}"}
        records.append(record)

    output = (args.out or root / "reports" / "INPUT_INVENTORY.md").resolve()
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(render(root, records, threshold_mb, sample_rows), encoding="utf-8")
    print(json.dumps({"output": str(output), "files": len(records)}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
